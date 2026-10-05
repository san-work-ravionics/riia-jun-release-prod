"""SYNTHETIC Console-style files for F42 P2 tests (fake symbols, ids and numbers only)."""
from __future__ import annotations

import csv
import io
from datetime import date

import openpyxl

TB_HEADER = ["symbol", "isin", "trade_date", "exchange", "segment", "series", "trade_type",
             "auction", "quantity", "price", "trade_id", "order_id", "order_execution_time",
             "expiry_date"]
LEDGER_HEADER = ["particulars", "posting_date", "cost_center", "voucher_type", "debit",
                 "credit", "net_balance"]
PNL_HEADER = ["Symbol", "ISIN", "Quantity", "Buy Value", "Sell Value", "Realized P&L",
              "Realized P&L Pct.", "Previous Closing Price", "Open Quantity",
              "Open Quantity Type", "Open Value", "Unrealized P&L", "Unrealized P&L Pct."]


def tb_row(symbol="NIFTY26OCT24000CE", d="2026-10-01", side="buy", qty=75, price=100.5,
           tid="T1", oid="O1", exp="2026-10-27"):
    return [symbol, "", d, "NFO", "FO", "OPT", side, "false", qty, price, tid, oid,
            f"{d}T10:15:00", exp]


def csv_bytes(header, rows, delimiter=",", bom=False) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=delimiter, lineterminator="\r\n")
    w.writerow(header)
    w.writerows(rows)
    return (("﻿" if bom else "") + buf.getvalue()).encode("utf-8")


def tradebook(rows=None, **kw) -> bytes:
    return csv_bytes(TB_HEADER, rows if rows is not None else [
        tb_row(tid="T1", oid="O1"),
        tb_row(symbol="NIFTY26OCT24000PE", tid="T2", oid="O2", side="sell", d="2026-10-02"),
        tb_row(symbol="BANKNIFTY26NOVFUT", tid="T3", oid="O3", d="2026-06-15", exp="2026-11-24"),
    ], **kw)


def ledger(rows=None) -> bytes:
    return csv_bytes(LEDGER_HEADER, rows if rows is not None else [
        ["Fake settlement A", "2026-09-01", "", "Book Voucher", "", "500.50", "1000"],
        ["Fake settlement A", "2026-09-01", "", "Book Voucher", "", "500.50", "1500.5"],
        ["Fake charge B", "2026-09-02", "", "Journal", "12", "", ""],
    ])


def pnl_xlsx(symbols=None, offset=3, period=("2026-04-01", "2026-10-05"), other=True,
             extra_sheet=False) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "F&O"
    if period:
        ws.append([None, f"P&L Statement for F&O from {period[0]} to {period[1]}"])
    for _ in range(offset - 1):
        ws.append([])
    ws.append([None, "Charges", 1234.5])
    ws.append([None, "Other Credit & Debit", "(50.25)"])
    ws.append([None, "Realized P&L", 9999.0])
    ws.append([None, "Unrealized P&L", -10])
    ws.append([None, "Brokerage", 20])
    ws.append([])
    ws.append(PNL_HEADER)
    for s in (symbols if symbols is not None else [
            ("NIFTY26OCT24000CE", 100.0), ("NIFTY26OCT24000PE", -40.5)]):
        ws.append([s[0], "", 75, 1000.0, 1100.0, s[1], 1.5, 10.0, 0, "", 0, 0.0, 0.0])
    if other:
        o = wb.create_sheet("Other Debits and Credits")
        o.append(["Date", "Particulars", "Amount"])
        o.append([date(2026, 5, 5), "Fake DP charge", -15.0])
    if extra_sheet:
        wb.create_sheet("Notes").append(["x"])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
